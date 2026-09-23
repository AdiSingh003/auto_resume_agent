import pymupdf
from docx import Document

from backend.tools.resume_parser import parse_resume, validate_upload


def test_validate_upload():
    assert validate_upload("cv.pdf", 1000)[0]
    assert validate_upload("cv.DOCX", 1000)[0]
    assert validate_upload("cv.txt", 1000)[0]
    assert not validate_upload("cv.doc", 1000)[0]
    assert not validate_upload("cv.exe", 1000)[0]
    assert not validate_upload("cv", 1000)[0]
    assert not validate_upload("cv.pdf", 0)[0]
    assert not validate_upload("cv.pdf", 11 * 1024 * 1024)[0]


def test_parse_docx(tmp_path):
    doc = Document()
    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("")
    doc.add_paragraph("Python developer")
    path = tmp_path / "cv.docx"
    doc.save(path)
    assert parse_resume(str(path)) == "Jane Doe\nPython developer"


def test_parse_txt(tmp_path):
    path = tmp_path / "cv.txt"
    path.write_text("Jane – Doe", encoding="utf-8")
    assert parse_resume(str(path)) == "Jane – Doe"


def test_parse_pdf(tmp_path):
    path = tmp_path / "cv.pdf"
    with pymupdf.open() as pdf:
        pdf.new_page().insert_text((72, 72), "Jane Doe Python Engineer")
        pdf.save(path)
    assert "Jane Doe Python Engineer" in parse_resume(str(path))
